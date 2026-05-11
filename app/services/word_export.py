"""Word 보고서 생성.

python-docx 사용. KICT 보고서 양식 기반.
한글 폰트는 맑은 고딕 사용. 차트 이미지 대신 표 중심 구성
(서버 환경 한글 폰트 의존성 회피).

구성:
1. 표지
2. 응답 현황 요약
3. 기술분류 분석 (표)
4. 연구기간 분석 (표)
5. 성과물 유형 분석 (표)
6. 제안자 소속 분포 (표)
7. 주요 키워드 (표)
8. 응답자 의견 발췌
9. 부록: 전체 응답 명단
"""

from datetime import datetime
from io import BytesIO

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.database import get_pool
from app.services import statistics


# ──────────────────────────────────────────────────────────────
# 스타일 헬퍼
# ──────────────────────────────────────────────────────────────


def _set_font(run, name: str = "맑은 고딕", size: int = 11,
              bold: bool = False, color: tuple[int, int, int] | None = None) -> None:
    """Run에 한글·영문 폰트 모두 설정."""
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    # 한글 폰트 별도 지정 (asia)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        from docx.oxml import OxmlElement
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    """한글 폰트 적용 제목 추가."""
    sizes = {1: 16, 2: 13, 3: 11}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    _set_font(run, size=sizes.get(level, 11), bold=True, color=(30, 64, 175))


def _add_para(doc: Document, text: str, size: int = 11, bold: bool = False,
              align=WD_ALIGN_PARAGRAPH.LEFT, color: tuple | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    _set_font(run, size=size, bold=bold, color=color)


def _set_cell_text(cell, text: str, bold: bool = False, size: int = 10,
                   align=WD_ALIGN_PARAGRAPH.LEFT,
                   bg_color: str | None = None) -> None:
    """표 셀에 텍스트와 스타일 설정."""
    cell.text = ""  # 기본 빈 단락 정리
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(str(text))
    _set_font(run, size=size, bold=bold, color=((255, 255, 255) if bg_color else None))
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    if bg_color:
        from docx.oxml import OxmlElement
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), bg_color)
        tcPr.append(shd)


def _build_stat_table(doc: Document, data: list[dict], label_key: str,
                      value_key: str = "count",
                      label_header: str = "구분") -> None:
    """표준 통계 표 생성: [구분 | 응답 수 | 비율(%)]."""
    total = sum(int(d.get(value_key, 0)) for d in data)

    table = doc.add_table(rows=1 + len(data) + 1, cols=3)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 헤더
    hdr = table.rows[0].cells
    _set_cell_text(hdr[0], label_header, bold=True,
                   align=WD_ALIGN_PARAGRAPH.CENTER, bg_color="1E40AF")
    _set_cell_text(hdr[1], "응답 수", bold=True,
                   align=WD_ALIGN_PARAGRAPH.CENTER, bg_color="1E40AF")
    _set_cell_text(hdr[2], "비율(%)", bold=True,
                   align=WD_ALIGN_PARAGRAPH.CENTER, bg_color="1E40AF")

    # 본문
    for i, d in enumerate(data, start=1):
        cnt = int(d.get(value_key, 0))
        pct = (cnt / total * 100) if total > 0 else 0
        row = table.rows[i].cells
        _set_cell_text(row[0], str(d.get(label_key, "")))
        _set_cell_text(row[1], str(cnt), align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_text(row[2], f"{pct:.1f}", align=WD_ALIGN_PARAGRAPH.CENTER)

    # 합계
    last = table.rows[-1].cells
    _set_cell_text(last[0], "합계", bold=True)
    _set_cell_text(last[1], str(total), bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_cell_text(last[2], "100.0" if total > 0 else "0.0", bold=True,
                   align=WD_ALIGN_PARAGRAPH.CENTER)


# ──────────────────────────────────────────────────────────────
# 섹션 빌더
# ──────────────────────────────────────────────────────────────


def _build_cover(doc: Document, total_count: int) -> None:
    """표지 페이지."""
    doc.add_paragraph().paragraph_format.space_after = Pt(60)

    _add_para(doc, "차세대 모듈러 건축 핵심기술 개발 기획연구",
              size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
              color=(100, 100, 100))

    doc.add_paragraph().paragraph_format.space_after = Pt(30)

    _add_para(doc, "기술 수요조사 결과 보고", size=22, bold=True,
              align=WD_ALIGN_PARAGRAPH.CENTER, color=(30, 64, 175))

    doc.add_paragraph().paragraph_format.space_after = Pt(120)

    _add_para(doc, f"총 응답 수: {total_count}건", size=14,
              align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_para(doc, f"보고서 생성일: {datetime.now().strftime('%Y년 %m월 %d일')}",
              size=12, align=WD_ALIGN_PARAGRAPH.CENTER, color=(100, 100, 100))

    doc.add_paragraph().paragraph_format.space_after = Pt(180)
    _add_para(doc, "한국건설기술연구원 (KICT)", size=12, bold=True,
              align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_para(doc, "건축연구본부 OSC 그룹", size=11,
              align=WD_ALIGN_PARAGRAPH.CENTER, color=(100, 100, 100))

    doc.add_page_break()


def _build_summary_section(doc: Document, responses: list[dict],
                           tech_cats: list[dict], output_types: list[dict]) -> None:
    """1. 응답 현황 요약."""
    _add_heading(doc, "1. 응답 현황 요약", level=1)

    total = len(responses)
    _add_para(doc, f"○ 총 응답 수: {total}건")

    if total == 0:
        _add_para(doc, "○ 현재 접수된 응답이 없습니다.", color=(180, 0, 0))
        return

    # 응답 기간
    dates = [r["created_at"] for r in responses if r.get("created_at")]
    if dates:
        first = min(dates).strftime("%Y-%m-%d")
        last = max(dates).strftime("%Y-%m-%d")
        _add_para(doc, f"○ 응답 기간: {first} ~ {last}")

    # 최다 응답
    if tech_cats:
        top = max(tech_cats, key=lambda x: x["count"])
        _add_para(doc,
                  f"○ 최다 응답 기술분류: {top['category']} ({top['count']}건, "
                  f"{top['count']/total*100:.1f}%)")

    if output_types:
        top = max(output_types, key=lambda x: x["count"])
        _add_para(doc,
                  f"○ 최다 응답 성과물유형: {top['output_type']} ({top['count']}건, "
                  f"{top['count']/total*100:.1f}%)")


def _build_tech_category_section(doc: Document, data: list[dict]) -> None:
    _add_heading(doc, "2. 기술 분류 분석", level=1)
    _add_para(doc, "○ 응답자가 제안한 기술의 분류별 분포 (중복 응답 가능)",
              size=10, color=(80, 80, 80))
    doc.add_paragraph()
    _build_stat_table(doc, data, label_key="category", label_header="기술분류")


def _build_research_period_section(doc: Document, data: list[dict]) -> None:
    _add_heading(doc, "3. 연구기간 분석", level=1)
    _add_para(doc, "○ 제안된 과제의 연구기간 분포", size=10, color=(80, 80, 80))
    doc.add_paragraph()
    formatted = [{"label": f"{d['years']}년", "count": d["count"]} for d in data]
    _build_stat_table(doc, formatted, label_key="label", label_header="연구기간")


def _build_output_type_section(doc: Document, data: list[dict]) -> None:
    _add_heading(doc, "4. 성과물 유형 분석", level=1)
    _add_para(doc, "○ 예상 최종 성과물 유형 분포 (중복 응답 가능)",
              size=10, color=(80, 80, 80))
    doc.add_paragraph()
    _build_stat_table(doc, data, label_key="output_type", label_header="성과물 유형")


def _build_organization_section(doc: Document, data: list[dict]) -> None:
    _add_heading(doc, "5. 제안자 소속 분포", level=1)
    _add_para(doc, "○ 산·학·연·관 분류 (자동 분류 기준이므로 별도 정성 검토 권장)",
              size=10, color=(80, 80, 80))
    doc.add_paragraph()
    _build_stat_table(doc, data, label_key="sector", label_header="구분")


def _build_keywords_section(doc: Document, data: list[dict], top_n: int = 20) -> None:
    _add_heading(doc, "6. 주요 키워드 (상위 20)", level=1)
    _add_para(doc, "○ 응답자가 입력한 키워드의 빈도 분석",
              size=10, color=(80, 80, 80))
    doc.add_paragraph()
    if not data:
        _add_para(doc, "  (키워드 입력 응답이 없습니다)", color=(180, 0, 0))
        return
    _build_stat_table(doc, data[:top_n], label_key="keyword", label_header="키워드")


def _build_opinion_excerpts(doc: Document, responses: list[dict]) -> None:
    """7. 응답자 의견 발췌 (기술 개요·필요성 일부)."""
    _add_heading(doc, "7. 주요 응답자 의견 발췌", level=1)
    _add_para(doc, "○ 응답자가 작성한 기술 개요 및 필요성 중 일부 발췌",
              size=10, color=(80, 80, 80))

    if not responses:
        return

    for i, r in enumerate(responses[:10], start=1):  # 최대 10건
        _add_heading(doc, f"7.{i} {r.get('project_title', '제목 없음')}", level=2)

        overview = (r.get("tech_overview") or "").strip()
        necessity = (r.get("tech_necessity") or "").strip()

        if overview:
            _add_para(doc, f"○ 기술 개요", bold=True, size=10)
            _add_para(doc, f"  {overview[:500]}{'...' if len(overview) > 500 else ''}",
                      size=10)

        if necessity:
            _add_para(doc, f"○ 기술 필요성", bold=True, size=10)
            _add_para(doc, f"  {necessity[:500]}{'...' if len(necessity) > 500 else ''}",
                      size=10)

        proposer_org = r.get("proposer_organization", "")
        if proposer_org:
            _add_para(doc, f"   - 제안: {proposer_org}",
                      size=9, color=(120, 120, 120))


def _build_appendix(doc: Document, responses: list[dict]) -> None:
    """8. 부록: 전체 응답 명단."""
    doc.add_page_break()
    _add_heading(doc, "부록. 전체 응답 명단", level=1)

    if not responses:
        _add_para(doc, "(응답 없음)")
        return

    table = doc.add_table(rows=1 + len(responses), cols=5)
    table.style = "Light Grid Accent 1"

    headers = ["No.", "과제명", "제안자", "소속", "기간"]
    for i, h in enumerate(headers):
        _set_cell_text(table.rows[0].cells[i], h, bold=True,
                       align=WD_ALIGN_PARAGRAPH.CENTER, bg_color="1E40AF")

    for i, r in enumerate(responses, start=1):
        row = table.rows[i].cells
        _set_cell_text(row[0], str(i), align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
        _set_cell_text(row[1], r.get("project_title", ""), size=9)
        _set_cell_text(row[2], r.get("proposer_name", ""), size=9)
        _set_cell_text(row[3], r.get("proposer_organization", ""), size=9)
        _set_cell_text(row[4], f"{r.get('research_period_years', '—')}년",
                       size=9, align=WD_ALIGN_PARAGRAPH.CENTER)


# ──────────────────────────────────────────────────────────────
# 데이터 조회
# ──────────────────────────────────────────────────────────────


async def _fetch_all_responses() -> list[dict]:
    pool = await get_pool()
    async with pool.connection() as conn:
        async with conn.cursor() as cur:
            await cur.execute("""
                SELECT * FROM responses ORDER BY created_at ASC;
            """)
            rows = await cur.fetchall()
            cols = [d[0] for d in cur.description]
    return [dict(zip(cols, row)) for row in rows]


# ──────────────────────────────────────────────────────────────
# 메인 엔트리
# ──────────────────────────────────────────────────────────────


async def build_report() -> bytes:
    """전체 Word 보고서를 bytes로 반환."""
    responses = await _fetch_all_responses()
    tech_cats = await statistics.aggregate_tech_categories()
    research_periods = await statistics.aggregate_research_periods()
    output_types = await statistics.aggregate_output_types()
    orgs = await statistics.aggregate_organizations()
    keywords = await statistics.aggregate_keywords_top(n=50)

    doc = Document()

    # 페이지 여백
    section = doc.sections[0]
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)

    # 기본 스타일
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(11)

    # 섹션 빌드
    _build_cover(doc, len(responses))
    _build_summary_section(doc, responses, tech_cats, output_types)
    _build_tech_category_section(doc, tech_cats)
    _build_research_period_section(doc, research_periods)
    _build_output_type_section(doc, output_types)
    _build_organization_section(doc, orgs)
    _build_keywords_section(doc, keywords)
    _build_opinion_excerpts(doc, responses)
    _build_appendix(doc, responses)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
